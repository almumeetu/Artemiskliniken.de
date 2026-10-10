"""Build the editable ARTEMIS sales documents from one content source.

Run with the bundled document runtime. Render each DOCX with render_docx.py;
the PDFs exported by that renderer are the client-facing counterparts.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = '143652'
INK = '243B4C'
MUTED = '526576'
TEAL = '087EA4'
PALE = 'F1F6F8'
LINE = 'D9D9D9'
DATE = '10.10.2026'
VALID = '24.10.2026'
OFFER = 'WSS ART 2026 1010'
DEMO = 'https://artemiskliniken-de.vercel.app/'
SOURCES = [
    ('ARTEMIS Hauptauftritt', 'https://www.artemiskliniken.de/'),
    ('ARTEMIS Leverkusen', 'https://www.artemiskliniken.de/standorte/artemis-augenzentrum-leverkusen/'),
    ('ARTEMIS Opladen', 'https://www.artemiskliniken.de/standorte/artemis-augenarzt-praxis-opladen/'),
]
TAX = 'Zuzüglich gesetzlicher Umsatzsteuer, sofern anwendbar. Die steuerliche Behandlung wird vor Beauftragung bestätigt.'


def run_format(run, size=10.2, bold=False, color=INK):
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), 'Arial')
    return run


def hyperlink(p, label, url, size=9):
    rid = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    prop = OxmlElement('w:rPr')
    f = OxmlElement('w:rFonts')
    f.set(qn('w:ascii'), 'Arial'); f.set(qn('w:hAnsi'), 'Arial')
    prop.append(f)
    c = OxmlElement('w:color'); c.set(qn('w:val'), TEAL); prop.append(c)
    s = OxmlElement('w:sz'); s.set(qn('w:val'), str(int(size * 2))); prop.append(s)
    r.append(prop)
    t = OxmlElement('w:t'); t.text = label; r.append(t)
    h.append(r); p._p.append(h)


def rich(p, text, size=10.2, color=INK, bold=False):
    for part in re.split(r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*)', text):
        if not part:
            continue
        m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
        if m:
            hyperlink(p, m[1], m[2], size)
        elif part.startswith('**'):
            run_format(p.add_run(part[2:-2]), size, True, color)
        else:
            run_format(p.add_run(part), size, bold, color)


class AgencyDoc:
    def __init__(self, package, name, title, cover=False, compact=False):
        self.package = package
        self.name = name
        self.doc = Document()
        self.md = []
        self.compact = compact
        self.size = 9.7 if compact else 10.2
        sec = self.doc.sections[0]
        sec.page_width = Cm(21); sec.page_height = Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(1.75)
        sec.top_margin = Cm(1.75 if compact else 1.85)
        sec.bottom_margin = Cm(1.6)
        sec.header_distance = Cm(.65); sec.footer_distance = Cm(.75)
        sec.different_first_page_header_footer = cover
        styles = self.doc.styles
        for name, size, before, after in [('Normal', self.size, 0, 7), ('Title', 27, 0, 10), ('Heading 1', 21, 0, 15), ('Heading 2', 11.5, 10, 5), ('Subtitle', 15, 0, 8)]:
            st = styles[name]
            st.font.name = 'Arial'; st.font.size = Pt(size)
            st.font.bold = name in ('Title', 'Heading 1', 'Heading 2')
            st.font.color.rgb = RGBColor.from_string('000000' if name != 'Normal' else INK)
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.line_spacing = 1.10
            st.paragraph_format.keep_with_next = name != 'Normal'
            prop = st._element.get_or_add_pPr()
            for tag in ['pBdr', 'shd']:
                node = prop.find(qn('w:' + tag))
                if node is not None: prop.remove(node)
            rpr = st._element.get_or_add_rPr()
            lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'de-DE'); rpr.append(lang)
            color = rpr.find(qn('w:color'))
            if color is not None:
                for attr in ('themeColor', 'themeTint', 'themeShade'):
                    color.attrib.pop(qn('w:' + attr), None)
        hp = sec.header.paragraphs[0]
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_TAB_ALIGNMENT.RIGHT)
        rich(hp, 'WEBDEV SOFTWARE SOLUTIONS\tARTEMIS Leverkusen und Opladen', 7.4, '000000')
        fp = sec.footer.paragraphs[0]
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_TAB_ALIGNMENT.RIGHT)
        rich(fp, OFFER + ' · ' + DATE + '\t', 7.4, MUTED)
        rich(fp, 'Seite ', 7.4, MUTED)
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
        fp._p.append(field)
        if cover:
            sec.first_page_header.paragraphs[0].text = ''
            first_footer = sec.first_page_footer.paragraphs[0]
            first_footer.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_TAB_ALIGNMENT.RIGHT)
            rich(first_footer, OFFER + '\t' + DATE, 8, MUTED)
        meta = self.doc.core_properties
        meta.title = title; meta.subject = 'Angebot zur Übernahme und Anpassung des regionalen Website-Konzepts'
        meta.author = 'Moyen Uddin | WebDev Software Solutions'
        meta.last_modified_by = 'WebDev Software Solutions'
        meta.comments = 'Projektunterlagen für ARTEMIS Leverkusen und Opladen'
        meta.created = meta.modified = datetime(2026, 10, 10, 0, 0, tzinfo=timezone.utc)
        meta.keywords = 'ARTEMIS, Leverkusen, Opladen, Projektangebot, WebDev Software Solutions'

    def p(self, text, size=None, color=INK, after=7, before=0, bold=False, keep=False, style=None):
        p = self.doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = keep
        rich(p, text, size or self.size, color, bold)
        self.md.append(text + '\n')
        return p

    def h(self, text, level=2):
        p = self.doc.add_paragraph(style='Heading ' + str(level))
        rich(p, text, 21 if level == 1 else 11.5, '000000', True)
        self.md.append('#' * level + ' ' + text + '\n')
        return p

    def page(self, title):
        self.doc.add_page_break()
        self.h(title, 1)

    def table(self, headers, rows, widths, size=9.1, align_right=(), highlight_last=False):
        table = self.doc.add_table(rows=0, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
        usable = 17.5
        widths = [usable * w / sum(widths) for w in widths]
        grid = table._tbl.tblGrid
        for child in list(grid): grid.remove(child)
        for w in widths:
            col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(int(Cm(w).twips))); grid.append(col)
        props = table._tbl.tblPr
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            b = OxmlElement('w:' + edge)
            for k,v in {'val':'single', 'sz':'5', 'color':LINE}.items(): b.set(qn('w:' + k),v)
            borders.append(b)
        props.append(borders)
        for ri, row in enumerate([headers] + rows):
            new_row = table.add_row()
            trpr = new_row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
            for ci, text in enumerate(row):
                cell = new_row.cells[ci]; cell.width = Cm(widths[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), NAVY if ri == 0 else (PALE if ri % 2 else 'FFFFFF'))
                if highlight_last and ri == len(rows): shd.set(qn('w:fill'), 'E1EEF2')
                cp.append(shd)
                mar = OxmlElement('w:tcMar')
                for side, value in [('top',95),('bottom',95),('start',110),('end',110)]:
                    m=OxmlElement('w:'+side); m.set(qn('w:w'),str(value));m.set(qn('w:type'),'dxa');mar.append(m)
                cp.append(mar)
                p=cell.paragraphs[0]
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.06
                if ci in align_right: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                rich(p, str(text), size, 'FFFFFF' if ri == 0 else INK, ri == 0 or (highlight_last and ri == len(rows)))
        spacer=self.doc.add_paragraph()
        spacer.paragraph_format.space_after=Pt(3)
        spacer.paragraph_format.line_spacing=Pt(2)
        self.md.append('| ' + ' | '.join(headers) + ' |\n| ' + ' | '.join(['---']*len(headers)) + ' |\n' + '\n'.join('| ' + ' | '.join(str(x).replace('\n',' ') for x in r) + ' |' for r in rows) + '\n')
        return table

    def picture(self, file, width, height=None, crop_bottom=0, alt=''):
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
        pic=p.add_run().add_picture(str(file),width=Cm(width),height=Cm(height) if height else None)
        pic._inline.docPr.set('descr',alt)
        if crop_bottom:
            fill = pic._inline.graphic.graphicData.pic.blipFill
            rect = OxmlElement('a:srcRect'); rect.set('b',str(crop_bottom))
            fill.insert(1,rect)
        return p

    def source_note(self):
        labels = ' · '.join(f'[{i} {title}]({url})' for i,(title,url) in enumerate(SOURCES,1))
        self.p(labels, 7.6, MUTED, after=3)
        self.p('Grundlage: öffentliche Text- und Linkinhalte, eingesehen am 10.10.2026. Die abgeleiteten Vorteile sind Projektziele; Nutzungsdaten, Ladezeiten und interne Systeme wurden nicht vergleichend bewertet.', 7.6, MUTED, after=0)

    def save(self, source=True):
        target=self.package/'02_EDITIERBARE_DOKUMENTE'/(self.name+'.docx')
        self.doc.save(target)
        if source:
            (self.package/'04_QUELLDATEIEN'/(self.name+'.md')).write_text('\n'.join(self.md),encoding='utf-8')
        return target


def proposal(package):
    d=AgencyDoc(package,'01_ARTEMIS_Projektangebot_DE','Regionale Website für ARTEMIS Leverkusen und Opladen',cover=True)
    assets=package/'04_QUELLDATEIEN'
    d.picture(assets/'06_WebDev_Software_Solutions_Logo.png',3.0,alt='WebDev Software Solutions')
    d.p('PROJEKTANGEBOT',8.5,TEAL,after=6,before=5,bold=True)
    d.p('Regionale Website\nfür ARTEMIS',29,'000000',after=4,bold=True,style='Title')
    d.p('Leverkusen und Opladen',16,'000000',after=11)
    d.p('Übernahme und Anpassung des vorliegenden Website-Konzepts\nmit Veröffentlichung und technischer Übergabe.',10.7,after=13)
    d.picture(assets/'07_BILDMATERIAL/01_Desktop_Startseite.png',17.5,9.72,80000,'Startseite des vorliegenden regionalen ARTEMIS-Website-Konzepts')
    d.p('Einblick in den vorliegenden Entwurf · Inhalte und Veröffentlichung nach ARTEMIS-Freigabe',7.6,MUTED,after=12)
    d.table(['Projektfestpreis','Finalisierung','Angebot gültig bis'],[['4.950 EUR','3–4 Wochen¹',VALID]],[1,1,1],size=11.8)
    d.p(TAX,7.6,MUTED,after=3)
    d.p('¹ Ab Beauftragung, Startzahlung und vollständigen Inhalten, Freigaben und Zugängen.',7.6,MUTED,after=11)
    d.p('Ihr Ansprechpartner  Moyen Uddin',9.2,bold=True,after=3)
    d.p('WebDev Software Solutions · Küppersteg, 51373 Leverkusen\n+49 172 9766016 · [info@webdevsoftwaresolutions.com](mailto:info@webdevsoftwaresolutions.com)\n[Website-Konzept ansehen](https://artemiskliniken-de.vercel.app/) · [webdevss.tech](https://webdevss.tech/)',8.6,after=0)

    d.page('Ihr regionales Websitekonzept')
    d.p('**Der Entwurf liegt bereits vor.** Wir bieten ARTEMIS an, das regionale Website-Konzept für Leverkusen und Opladen zu übernehmen, gemeinsam anzupassen und nach Freigabe zu veröffentlichen. Die Vorschau macht Gestaltung und Patientenwege schon vor einer Beauftragung besprechbar.')
    d.h('Was sich für Ihre Standorte verbessern lässt')
    d.p('Der nationale ARTEMIS-Auftritt bietet eine Standortsuche und medizinische Fachinformationen. Die beiden Standortseiten enthalten bereits lokale Angaben. [1–3] Daraus leiten wir folgende Ansatzpunkte für den regionalen Entwurf ab:',size=9.7)
    d.table(['Beobachtung am bestehenden Auftritt','Umsetzung im regionalen Konzept'],[
        ['Leverkusen und Opladen werden auf eigenen Standortseiten dargestellt. [2, 3]','Ein gemeinsamer Standortvergleich ordnet Adresse, Sprechzeiten, Leistungen und Kontakt direkt zu.'],
        ['Leverkusen verlinkt eine Onlinebuchung; Opladen nennt Telefon und Kontaktanfrage. [2, 3]','Die Terminansicht erklärt den jeweiligen Weg und führt gezielt zum bestehenden Buchungsdienst oder zur Praxis.'],
        ['Die Hauptnavigation erschließt ein breites medizinisches Angebot. [1]','Der regionale Einstieg priorisiert Standortwahl, Behandlungen, Team und den nächsten Kontakt.'],
    ],[.47,.53],size=9)
    d.h('Nutzen für Patienten und Management')
    d.p('Patientinnen und Patienten sollen schneller erkennen, welche Praxis für ihr Anliegen infrage kommt und wie sie diese erreichen. Gebündelte Kontaktangaben können Rückfragen zu Erreichbarkeit und Anfahrt verringern. Diese Wirkungen sind nach Einführung anhand abgestimmter Kriterien zu bewerten.',size=9.8)
    d.p('Sie erhalten einen konkreten, anpassbaren Entwurf mit Festpreis, zwei Korrekturrunden und dokumentierter Abnahme. Seitenumfang, technische Umsetzung und Übergabe sind in diesem Angebot festgelegt.',size=9.8)
    d.h('Einbindung in den ARTEMIS Auftritt')
    d.p('Kalkuliert ist eine eigenständige regionale Website auf einem von ARTEMIS freigegebenen Host. Marke, Fachinhalte und bestehende Terminprozesse bilden die Grundlage. Domain, Verlinkung und überlappende Inhalte stimmen wir vor Projektstart ab. Eine Integration in das zentrale CMS wird gesondert bewertet.',size=9.6)
    d.source_note()

    d.page('Patientenwege mit klarem Standortbezug')
    d.p('Der Entwurf ordnet Informationen entlang typischer Patientenfragen. Jeder Einstieg führt zu einem passenden nächsten Schritt; Akuthinweise bleiben unabhängig vom regulären Terminweg erreichbar.')
    d.table(['Station','Frage des Patienten','Vorgesehene Nutzerführung'],[
        ['Standort','Welche Praxis passt?','Leverkusen und Opladen mit Adresse, Leistungsprofil, Sprechzeiten und Anfahrt unterscheiden.'],
        ['Behandlung','Was ist für mich relevant?','Freigegebene Informationen zu Behandlungen, Diagnostik und Augenkrankheiten dem Standort zuordnen.'],
        ['Ärzteteam','Wer behandelt mich?','Bestätigte Profile mit Rolle und Standortzuordnung anzeigen.'],
        ['Kontakt','Wie erreiche ich die Praxis?','Telefon, E-Mail, Sprechzeiten und Routenlink im Standortkontext bereitstellen.'],
        ['Termin','Wie vereinbare ich einen Termin?','Für Leverkusen den bestätigten externen Buchungslink, für Opladen den telefonischen Weg anbieten.'],
        ['Akutfall','Wo finde ich Hinweise?','Freigegebene Akuthinweise von jeder Seite direkt zugänglich machen.'],
    ],[.17,.27,.56],size=9.4)
    d.h('Terminwege verständlich erklären')
    d.p('Die Terminansicht führt zum freigegebenen ARTEMIS-Buchungsdienst oder zur Telefonnummer der Praxis. Sie erfasst selbst keine Kontakt- oder Gesundheitsdaten und bestätigt keine Termine. Der Buchungsabschluss erfolgt im bestehenden Prozess.')
    d.h('Auf dem Smartphone direkt handeln')
    d.p('Mobile Navigation, Suche und Kontaktaktionen werden so finalisiert, dass Standortwahl, Anruf und Terminweg mit gut lesbaren Informationen und ausreichend großen Bedienelementen nutzbar sind. Die Abnahme umfasst Smartphone, Tablet und Desktop.')
    d.h('Akuthinweise separat erreichbar halten')
    d.p('ARTEMIS gibt Texte, Rufnummern und Zuständigkeiten für akute Anliegen frei. Die Website nimmt keine medizinische Ersteinschätzung vor. Der Zugang zu diesen Hinweisen setzt keine vorherige Auswahl von Standort, Behandlung oder Arzt voraus.')

    d.page('Leistungsumfang und Anpassungen')
    d.p('Der Festpreis umfasst die Übernahme des vorliegenden Konzepts als Projektbasis und dessen Finalisierung für zwei Standorte. Enthalten sind bis zu **35 Seitenansichten in deutscher Sprache**. Die verbindliche Seitenliste bestätigen wir in Woche 1.')
    d.table(['Seitenbereich','Enthaltene Ansichten','Anzahl'],[
        ['Regionaler Einstieg','Startseite','1'],
        ['Standorte','Übersicht und zwei Standortdetails','3'],
        ['Ärzteteam','Übersicht und bis zu fünf Profile','6'],
        ['Behandlungen','Übersicht und sechs Themen','7'],
        ['Diagnostik','Übersicht und drei Themen','4'],
        ['Augenkrankheiten','Übersicht und neun kurze Informationen','10'],
        ['Patienten und Akutfälle','Patienteninformation und Akuthinweise','2'],
        ['Rechtliche Informationen','Impressum und Datenschutz nach Freigabe','2'],
        ['Gesamt','Abschließend bestätigte Seitenliste','35'],
    ],[.30,.60,.10],size=9.2,align_right=(2,),highlight_last=True)
    d.h('Gestaltung und Inhalte abstimmen')
    d.p('Wir passen den Entwurf an die freigegebenen ARTEMIS-Vorgaben an, bearbeiten vorhandene deutsche Texte redaktionell und integrieren bestätigte Standortdaten sowie freigegebenes Bildmaterial. **Zwei gebündelte Korrekturrunden** innerhalb des vereinbarten Umfangs sind enthalten.')
    d.h('Funktionen finalisieren')
    d.p('Seitensuche, mobile Navigation, standortbezogene Kontaktkarten und Terminansicht gehören zum Umfang. Telefon-, E-Mail-, Routen- und Buchungslinks führen zu bestätigten Zielen. Externe Karten werden erst nach aktiver Auswahl geladen.')
    d.h('Änderungswünsche planbar umsetzen')
    d.p('Sie können Texte, Bildauswahl, Kontaktangaben und Darstellungsdetails im abgestimmten Umfang ändern lassen. Zusätzliche Seiten, Funktionen oder Integrationen bewerten wir vor Umsetzung und halten Auswirkungen auf Preis und Termin schriftlich fest.')

    d.page('Technische Qualität und Betrieb')
    d.h('Responsive Umsetzung und Browser')
    d.p('Die vorhandene Basis nutzt React, TypeScript und Vite. Wir prüfen die vereinbarten Seiten bei 320, 375, 768 und 1440 Pixeln Breite sowie mit vergrößertem Text. Zur Abnahme berücksichtigen wir aktuelle stabile Versionen von Chrome, Edge, Firefox und Safari sowie Safari unter iOS und Chrome unter Android. Versionen und Geräte oder Simulationen werden dokumentiert.')
    d.h('Zugänglichkeit im vereinbarten Umfang')
    d.p('Lesbare Schrift, geeignete Kontraste, Tastaturbedienung, sichtbarer Fokus, Dialogbedienung, Textvergrößerung und reduzierte Bewegung gehören zur Umsetzung und Prüfung. Die konkreten Prüfergebnisse werden mit den Abnahmeunterlagen übergeben.')
    d.h('Technische SEO und Seitenaufruf')
    d.p('Wir finalisieren Seitenpfade, Titel, Beschreibungen, Canonicals, strukturierte Standortdaten, Sitemap und Robots-Regeln für die freigegebene Domain. Die Konzeptbasis rendert Inhalte im Browser. Abrufbarkeit und Indexierbarkeit werden vor Veröffentlichung im vereinbarten Host geprüft und angepasst; überlappende Inhalte stimmen wir mit dem nationalen Auftritt ab.')
    d.h('Sicherheit und Datenverarbeitung')
    d.p('TLS, Sicherheitsheader, Abhängigkeiten und externe Ressourcen werden im Zielsystem geprüft. Zugänge werden auf die erforderlichen Aufgaben beschränkt. Der Umfang enthält keine eigene Patientenverwaltung und keine Speicherung von Terminanfragen. ARTEMIS bestätigt Rechts- und Datenschutzhinweise für die eingesetzten Dienste. Tracking erfolgt nur nach gesonderter Abstimmung.')
    d.h('Ladeverhalten und Veröffentlichung')
    d.p('Bildgrößen, verzögertes Laden geeigneter Inhalte, Codeaufteilung und Cache-Konfiguration werden auf den regionalen Auftritt abgestimmt. Wir prüfen auffälliges Ladeverhalten und Kerninteraktionen unter dokumentierten Bedingungen. Messwerte werden erst nach einer tatsächlichen Messung angegeben.')
    d.h('Leistungen außerhalb des Festpreises')
    d.p('Ein nationaler Relaunch, CMS-Neuentwicklung oder CMS-Migration, Patientenportal, eigene Formularverarbeitung, Buchungsbackend, Übersetzungen und laufende Kampagnen sind gesonderte Leistungen. Hosting, Domain, externe Lizenzen und kostenpflichtige Dienste werden separat getragen. Formale Barrierefreiheitszertifizierung, Penetrationstest sowie medizinische oder rechtliche Beratung sind nicht enthalten.',size=9.5)

    d.page('Projektablauf und technische Übergabe')
    d.p('Wir planen **3–4 Wochen für Anpassung und Finalisierung**. Der Zeitraum beginnt nach schriftlicher Beauftragung, Startzahlung sowie vollständiger Bereitstellung der erforderlichen Inhalte, Freigaben und Zugänge. Wartezeiten auf Entscheidungen oder externe Dienstleister verschieben die betroffenen Termine.')
    d.table(['Phase','Zeitraum','Arbeitsergebnis und Freigabe'],[
        ['Projektstart','Woche 1','Konzept gemeinsam durchgehen; Seiteninventar, Host, Terminziele, Ansprechpartner und Freigabeweg bestätigen.'],
        ['Gestaltung und Inhalte','Woche 1–2','Entwurf und deutsche Inhalte anpassen. Erste gebündelte Korrekturrunde; Design- und Inhaltsfreigabe durch ARTEMIS.'],
        ['Finalisierung','Woche 2–3','Frontend, Kontakt- und Terminwege, Domainkonfiguration und SEO-Grundlagen abschließen. Zweite Korrekturrunde auf der Prüfversion.'],
        ['Abnahme und Übergabe','Woche 3–4','Prüfergebnisse dokumentieren, vereinbarte Abweichungen beheben und Abnahme einholen. Nach schriftlicher Freigabe veröffentlichen.'],
    ],[.22,.15,.63],size=9.3)
    d.h('Website und Qualitätsnachweise')
    d.p('Sie erhalten den freigegebenen deutschen Auftritt für Leverkusen und Opladen mit den vereinbarten Seiten und Funktionen. Dazu gehören Seiten- und Linkinventar, dokumentierte Browser- und Bedienprüfungen sowie eine abgestimmte Liste etwaiger Restpunkte.')
    d.h('Projektdateien und Weiterbearbeitung')
    d.p('Wir übergeben den Projektquellcode, verwendbare Projektdateien und freigegebene Assets sowie Hinweise zu Build, Deployment und Konfiguration. Damit stehen die technischen Grundlagen für spätere Anpassungen zur Verfügung. Die Nutzungsrechte werden vor Beauftragung schriftlich festgelegt; Lizenzbedingungen für Drittanbieter-Komponenten, Bilder und Marken bleiben bestehen.')
    d.h('Veröffentlichung und Einweisung')
    d.p('Enthalten sind ein Deployment im vereinbarten Zielsystem, eine abschließende Funktionskontrolle und ein Übergabegespräch. Zuständigkeiten für Betrieb, Sicherung und Rückkehr zum vorherigen Stand werden vor Veröffentlichung zugeordnet.')

    d.page('Qualitätssicherung und Zusammenarbeit')
    d.p('Die Abnahme erfolgt anhand vereinbarter Kriterien mit freigegebenen Inhalten in der vorgesehenen Produktionsumgebung. Sie erhalten nachvollziehbare Prüfergebnisse und eine klare Grundlage für die Freigabe.')
    d.table(['Prüfbereich','Abnahmekriterium'],[
        ['Seiten und Inhalte','Alle vereinbarten Seiten sind direkt aufrufbar. Standortdaten, Leistungszuordnung, Arztrollen, Bilder und Rechtstexte entsprechen der ARTEMIS-Freigabe.'],
        ['Patientenwege','Navigation, Suche und Terminansicht sind bedienbar. Telefon-, E-Mail-, Routen- und Buchungslinks führen zum richtigen Standort. Akuthinweise sind direkt erreichbar.'],
        ['Darstellung und Bedienung','Keine abgeschnittenen Kerninhalte oder unbeabsichtigte horizontale Seitenbewegung in der Browsermatrix. Tastatur, Fokus und Dialoge sind geprüft; Text bleibt bei 200 % lesbar.'],
        ['Technischer Betrieb','Seitenpfade, Metadaten, Sitemap und Canonicals passen zum Zielsystem. Build, TLS, Sicherheitsheader und vereinbarte externe Dienste sind geprüft und dokumentiert.'],
    ],[.25,.75],size=9.2)
    d.p('Nicht erreichbare Kernseiten, falsche Standort- oder Terminziele und Bedienfehler, die Kernaufgaben verhindern, werden vor Abnahme behoben. Andere Restpunkte erhalten eine abgestimmte Zuständigkeit und Frist. ARTEMIS erteilt die schriftliche Abnahme und Freigabe zur Veröffentlichung.',size=9.7)
    d.h('Verantwortlichkeiten')
    d.table(['Beteiligte','Beitrag zum Projekt'],[
        ['ARTEMIS','Zentrale Ansprechperson benennen; Standort- und medizinische Angaben, Marke, Bildrechte, Rechts- und Datenschutzhinweise freigeben; Terminziele und Zugänge bereitstellen; Abnahme erteilen.'],
        ['WebDev Software Solutions','Vereinbarten Umfang umsetzen, Rückfragen und Korrekturen bündeln, Prüfungen dokumentieren, reproduzierbare Abweichungen im Umfang beheben und Website sowie Projektdateien übergeben.'],
        ['Gemeinsam','Seitenliste, Zielsystem, Freigabetermine und Betriebszuständigkeiten festlegen. Umfangsänderungen mit Preis- und Terminfolgen vor Umsetzung schriftlich vereinbaren.'],
    ],[.25,.75],size=9.1)

    d.page('Investition und nächste Schritte')
    d.p('Der Festpreis umfasst das vorliegende Website-Konzept als Projektbasis, die vereinbarten Anpassungen und die Finalisierung bis zur technischen Übergabe. Die Kalkulation gilt für den beschriebenen regionalen Umfang.',size=9.7)
    items=[
        ['Projektklärung und Inhaltsinventar','450 EUR'],
        ['Anpassung von Webdesign und Nutzerführung','1.100 EUR'],
        ['Deutsche Inhalte und Standortdaten','750 EUR'],
        ['Finalisierung von Frontend und Terminwegen','1.100 EUR'],
        ['Seitenstruktur und technische SEO','650 EUR'],
        ['Qualitätssicherung und Abnahme','550 EUR'],
        ['Veröffentlichung und Übergabe','350 EUR'],
        ['Projektgesamtpreis','4.950 EUR'],
    ]
    d.table(['Leistung','Betrag'],items,[.78,.22],size=9.4,align_right=(1,),highlight_last=True)
    d.p(TAX+' Vollständige Vertrags- und Rechnungsdaten werden vor Beauftragung schriftlich bestätigt.',8.2,MUTED,after=7)
    d.h('Zahlungsmeilensteine')
    d.table(['Meilenstein','Anteil','Betrag'],[
        ['Bei Beauftragung','30 %','1.485 EUR'],
        ['Nach Design- und Inhaltsfreigabe','40 %','1.980 EUR'],
        ['Nach Abnahme und vor Veröffentlichung','30 %','1.485 EUR'],
    ],[.62,.14,.24],size=9.2,align_right=(1,2))
    d.h('Optionale Betreuung')
    d.p('Betreuung kann für **95 EUR pro Monat** separat vereinbart werden. Wartung, kleinere Inhaltsänderungen, Stundenkontingent, Reaktionszeiten, Laufzeit und Betriebszuständigkeit werden schriftlich festgelegt. Externe Kosten und eine Rund-um-die-Uhr-Bereitschaft sind nicht enthalten.',size=9.3)
    d.h('Gemeinsam den Entwurf durchgehen')
    d.p('Gern zeigen wir Ihnen die Website in einem **30-minütigen Gespräch**. Sie geben Rückmeldung zum Entwurf; gemeinsam bestätigen wir Anpassungsbedarf, Zielsystem und Freigabeweg. Anschließend können Sie über Beauftragung und Projektstart entscheiden.',size=9.5)
    d.p('Angebot gültig bis **24.10.2026**.\nMoyen Uddin · WebDev Software Solutions\n[info@webdevsoftwaresolutions.com](mailto:info@webdevsoftwaresolutions.com) · +49 172 9766016',9.2,after=0)
    return d.save()


def comparison(package):
    d=AgencyDoc(package,'02_ARTEMIS_Managementvergleich_DE','Managementvergleich für ARTEMIS',compact=True)
    d.p('Managementvergleich',25,'000000',after=4,bold=True,style='Title')
    d.p('Nationaler ARTEMIS Auftritt und regionales Websitekonzept',11,'000000',after=5)
    d.p('Leverkusen und Opladen',9.3,MUTED,after=12)
    d.p('**Empfehlung:** Den vorliegenden Entwurf gemeinsam bewerten und die Anpassungen für einen regionalen Pilot festlegen. Website, Patientenwege und Gestaltung sind anhand des vorhandenen Konzepts besprechbar.',size=9.7)
    d.table(['Entscheidungsaspekt','Bestehender Auftritt','Regionaler Entwurf und Lieferziel'],[
        ['Patientenorientierung','Bundesweite Standortsuche und breites Informationsangebot. [1]','Einstieg für zwei Standorte mit klarer Zuordnung des nächsten Schritts.'],
        ['Lokale Informationen','Eigene Standortseiten mit Kontakt, Zeiten und Leistungen. [2, 3]','Gemeinsamer Standortvergleich und einheitliche Kontaktkarten.'],
        ['Terminweg','Leverkusen mit Onlinebuchung; Opladen mit Telefon und Kontaktanfrage. [2, 3]','Standortbezogene Erklärung und Weiterleitung in die bestehenden Prozesse.'],
        ['Mobile Bedienung','Keine vergleichende Geräte- oder Leistungsmessung durchgeführt.','Kompakte Navigation und Kontaktaktionen; dokumentierte mobile Abnahme.'],
        ['Management und Betrieb','Nationale Plattform für Marke und Standortnetz. [1]','Begrenzter Projektumfang, freigegebene Inhalte und dokumentierte technische Übergabe.'],
        ['Technische Grundlage','Interne Architektur und Qualität nicht bewertet.','Abgestimmte Metadaten, Seitenpfade, SEO-Grundlagen sowie Bedien- und Betriebsprüfung.'],
    ],[.22,.36,.42],size=8.7)
    d.h('Vorhandener Entwurf und klarer Angebotsrahmen')
    d.p('**4.950 EUR** für Konzeptübernahme, vereinbarte Anpassungen und Finalisierung. **3–4 Wochen** ab vollständigen Startvoraussetzungen; zwei Korrekturrunden und Quellcodeübergabe enthalten. '+TAX,size=9.2)
    d.p('Der Pilot ergänzt die nationale Plattform. Veröffentlichung und Inhalte benötigen ARTEMIS-Freigabe. Eine zentrale CMS-Integration wird separat bewertet. Ob sich Orientierung oder Kontaktaufnahme verbessern, ist nach Einführung zu prüfen.',size=9.2)
    d.p('**Nächster Schritt:** Den Entwurf in 30 Minuten vorstellen, Änderungsbedarf und Zielsystem bestätigen. [Projektvorschau öffnen]('+DEMO+') · Angebot gültig bis '+VALID+'.',size=9.2)
    d.source_note()
    return d.save()


def overview(package):
    d=AgencyDoc(package,'03_ARTEMIS_Projektuebersicht_DE','Projektübersicht für ARTEMIS',compact=True)
    d.p('Projektübersicht',25,'000000',after=4,bold=True,style='Title')
    d.p('Regionale Website für Leverkusen und Opladen',13,'000000',after=10)
    d.p('**Die Website als Konzept liegt vor.** WebDev Software Solutions bietet ARTEMIS die Übernahme, Anpassung und Finalisierung des regionalen Entwurfs an. Gestaltung, Standortinformationen und Kontaktwege können Sie vor Beauftragung konkret beurteilen.',size=10)
    d.h('Was Patienten vor Ort erhalten sollen')
    d.p('Einen gemeinsamen Einstieg zu den beiden Standorten, verständliche Informationen zu Leistungen und Team sowie passende Kontakt- und Terminwege. Akuthinweise bleiben unabhängig von der regulären Terminvereinbarung erreichbar.',size=9.7)
    d.h('Was im Projekt enthalten ist')
    d.table(['Bereich','Vereinbarter Rahmen'],[
        ['Inhalte und Seiten','Bis zu 35 deutsche Ansichten; Standorte, Behandlungen, Diagnostik, Team, Patienten- und Rechtstexte nach Freigabe.'],
        ['Anpassungen','Gestaltung und vorhandene Inhalte an ARTEMIS-Vorgaben anpassen; zwei gebündelte Korrekturrunden.'],
        ['Funktionen','Seitensuche, mobile Navigation, Kontaktkarten, Terminverweise und direkt erreichbare Akuthinweise.'],
        ['Qualität und Übergabe','Responsive Finalisierung, technische SEO, dokumentierte Abnahme, ein Deployment, Projektquellcode und Übergabegespräch.'],
    ],[.24,.76],size=9.2)
    d.table(['Projektpreis','Finalisierung','Zahlungsplan'],[['4.950 EUR','3–4 Wochen¹','30 / 40 / 30 %\n1.485 / 1.980 / 1.485 EUR']],[.27,.30,.43],size=9.6)
    d.p(TAX+' ¹ Nach schriftlicher Beauftragung, Startzahlung und vollständiger Bereitstellung der Inhalte, Freigaben und Zugänge.',8.1,MUTED)
    d.h('Einbindung und nächster Schritt')
    d.p('Der regionale Auftritt ergänzt die nationale Website und nutzt die bestehenden Terminprozesse. Kalkuliert ist ein freigegebener Host; eine Integration in das zentrale CMS wird gesondert bewertet. Inhalte, Markenmaterial und Veröffentlichung benötigen ARTEMIS-Freigabe.',size=9.5)
    d.p('**30-minütige Vorstellung vereinbaren:** Entwurf gemeinsam durchgehen, Anpassungsbedarf und Zielsystem bestätigen. [Website-Konzept ansehen]('+DEMO+').',size=9.8)
    d.p('Angebot '+OFFER+' · '+DATE+' · gültig bis '+VALID+'\nMoyen Uddin · WebDev Software Solutions\n[info@webdevsoftwaresolutions.com](mailto:info@webdevsoftwaresolutions.com) · +49 172 9766016',8.7,MUTED,after=0)
    return d.save()


def emails(package):
    signature='Mit freundlichen Grüßen\nMoyen Uddin\nWebDev Software Solutions\nKüppersteg, 51373 Leverkusen, NRW, Deutschland\n+49 172 9766016\ninfo@webdevsoftwaresolutions.com\nhttps://webdevss.tech/'
    messages=[
        ('A Erstkontakt','Website-Konzept für ARTEMIS Leverkusen und Opladen',[
            'Sehr geehrte Damen und Herren,',
            'bei der Durchsicht Ihres Webauftritts haben wir Ansatzpunkte gesehen, die Standortwahl und die Kontaktwege für Leverkusen und Opladen gemeinsam darzustellen. Darauf aufbauend haben wir ein regionales Website-Konzept ausgearbeitet, das wir Ihnen gern vorstellen möchten.',
            'Der Entwurf bündelt Standortinformationen, Leistungen und Ärzteteam. Die Terminansicht führt je nach Praxis zur bestehenden Onlinebuchung oder zum telefonischen Kontakt. Die nationale ARTEMIS-Plattform bleibt die Grundlage.',
            'Sie können den vorliegenden Entwurf hier ansehen: '+DEMO+' Er dient der Abstimmung; Inhalte und Veröffentlichung benötigen Ihre Freigabe.',
            'WebDev Software Solutions bietet die Anpassung an Ihre Vorgaben sowie die technische Finalisierung und Übergabe an. Gern zeigen wir Ihnen die Website in einem 30-minütigen Gespräch und besprechen Ihre Änderungswünsche.',
            'Passt Ihnen ein Termin in der kommenden Woche? Falls eine andere Person für den regionalen Webauftritt zuständig ist, freue ich mich über einen entsprechenden Hinweis.',
            signature,
        ],'Vor Versand Ansprechperson und Vorschau prüfen. Ohne Anlagen versendbar; bei Bedarf den einseitigen Managementvergleich beifügen.'),
        ('B Angebot nach dem Gespräch','Ihr Projektangebot für die regionale ARTEMIS-Website',[
            'Sehr geehrte Damen und Herren,',
            'vielen Dank für das Gespräch über den regionalen Webauftritt für Leverkusen und Opladen. Anbei erhalten Sie unser Projektangebot und den kompakten Managementvergleich.',
            'Das Angebot umfasst die Übernahme des vorliegenden Website-Konzepts als Projektbasis, die vereinbarten Anpassungen und die Finalisierung bis zur Veröffentlichung und technischen Übergabe. Zwei gebündelte Korrekturrunden und die Übergabe des Projektquellcodes sind enthalten.',
            'Der Projektpreis beträgt 4.950 EUR zuzüglich gesetzlicher Umsatzsteuer, sofern anwendbar. Die steuerliche Behandlung bestätigen wir vor Beauftragung. Für Anpassung und Finalisierung planen wir 3–4 Wochen ab vollständigen Startvoraussetzungen.',
            'Die Projektvorschau finden Sie unter '+DEMO+' Als nächsten Schritt bestätigen wir gemeinsam den Anpassungsbedarf, das Zielsystem, die Ansprechpartner und den Freigabeweg. Anschließend können wir Beauftragung und Projektstart schriftlich festhalten.',
            'Das Angebot '+OFFER+' ist bis zum '+VALID+' gültig. Gern erläutere ich Ihnen offene Punkte.',
            signature,
        ],'Nur nach einem tatsächlich geführten Gespräch verwenden. Anlagen: 01_ARTEMIS_Projektangebot_DE.pdf und 02_ARTEMIS_Managementvergleich_DE.pdf. Angaben zum Gespräch, Angebotsdatum und Gültigkeit vor Versand prüfen.'),
        ('C Rückfrage nach fünf Arbeitstagen','Rückfrage zum Website-Angebot Leverkusen und Opladen',[
            'Sehr geehrte Damen und Herren,',
            'ich möchte kurz an unser Projektangebot für Leverkusen und Opladen anknüpfen. Konnten Sie den Entwurf und die Unterlagen bereits besprechen?',
            'Gern gehe ich mit Ihnen mögliche Anpassungen an Gestaltung, Inhalten oder der technischen Einbindung durch. Wenn Ihre interne Abstimmung noch Zeit benötigt, genügt mir eine kurze Rückmeldung zum passenden Zeitpunkt.',
            'Ist ein kurzer Austausch in der kommenden Woche für Sie sinnvoll?',
            signature,
        ],'Fünf Arbeitstage nach Versand des Angebots im bestehenden E-Mail-Verlauf senden. Bei abgelaufener Gültigkeit zuerst das Angebot aktualisieren. Keine automatische Nachfassserie.'),
    ]
    d=AgencyDoc(package,'05_ARTEMIS_Emailvorlagen_DE','E Mail Vorlagen für ARTEMIS')
    text=['E-MAIL-VORLAGEN | ARTEMIS LEVERKUSEN UND OPLADEN','WebDev Software Solutions','Nur die jeweilige E-Mail kopieren; interne Versandhinweise nicht mitsenden.','']
    for idx,(heading,subject,paragraphs,note) in enumerate(messages):
        if idx: d.doc.add_page_break()
        d.p('E Mail Vorlagen für ARTEMIS' if idx==0 else 'ARTEMIS Leverkusen und Opladen',23 if idx==0 else 18,'000000',after=7,bold=True,style='Title')
        d.h(heading)
        d.p('**Betreff:** '+subject,10,after=12)
        for p in paragraphs:
            d.p(p,10.2,after=8)
        d.p('Interner Versandhinweis',8.5,'000000',before=7,after=3,bold=True)
        d.p(note,8.3,MUTED)
        text.extend([heading,'Betreff: '+subject,'','\n\n'.join(paragraphs),'','VERSANDHINWEIS (nicht mitsenden): '+note,'','-'*60,''])
    (package/'03_KOMMUNIKATION_UND_DEMO/01_Emailvorlagen_DE.txt').write_text('\n'.join(text),encoding='utf-8')
    return d.save(source=False)


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--package',type=Path,required=True)
    package=ap.parse_args().package.resolve()
    for fn in [proposal,comparison,overview,emails]:
        print(fn(package))
    data_path=package/'04_QUELLDATEIEN/05_Angebotsdaten.json'
    data=json.loads(data_path.read_text())
    data['version']='5'
    data['commercial_basis']='Übernahme des vorliegenden Website-Konzepts als Projektbasis mit vereinbarten Anpassungen, Finalisierung, Veröffentlichung und technischer Übergabe'
    data['review_date']=DATE
    data['correction_rounds']=2
    data['items'][1][0]='Anpassung von Webdesign und Nutzerführung'
    data['items'][3][0]='Finalisierung von Frontend und Terminwegen'
    data_path.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


if __name__=='__main__':
    main()
